#!/usr/bin/env python
"""Generate two Word documents for the revision:
   1. TREND_figure_legends.docx     — detailed legends + a 'how each figure is plotted' section
   2. TREND_figure_conclusions.docx — the conclusion drawn from each figure (kept separate)
Run after make_figures.py. Dep: python-docx  (pip install python-docx)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent

# ----------------------------------------------------------------- content
LEGENDS = [
("Figure 1. Sequence-variant activity landscapes for representative transcription-factor motifs.",
 "Each of the twelve panels corresponds to a single transcription-factor position-weight matrix (PWM) "
 "from the ENCODE/MotifDb motif collection; the panel title gives the transcription factor, the source "
 "database, and the motif accession. Within a panel, every plotted point is one synthetic enhancer "
 "variant — a tandem-repeat reporter built from a distinct binding-site sequence sampled from that PWM. "
 "Each variant is positioned by its measured activity in OV8 ovarian-cancer cells (x-axis) and in IOSE "
 "normal ovarian surface-epithelial cells (y-axis), expressed as the mean RNA-to-DNA barcode ratio "
 "(reporter activity) on linear axes that are shared across all panels. The fill colour of each point "
 "encodes the variant's relative binding affinity for its own PWM, computed as the fraction-of-consensus "
 "score (the product across motif positions of the base probability divided by that column's maximum "
 "probability; a value of 1.0 corresponds to the consensus, i.e. optimal, sequence) and shown on the "
 "viridis colour scale at right. The dashed diagonal (y = x) denotes equal activity in tumour and normal "
 "cells (no selectivity); points below the diagonal are tumour-selective and points above are "
 "normal-selective. In every panel the consensus variant — operationally defined as the surviving variant "
 "with the highest relative affinity — is drawn as a thick black-outlined dot, and the most tumour-selective "
 "variant — the variant with the highest OV8/IOSE ratio among those with OV8 activity ≥ 0.1 — is drawn "
 "as a red star. The in-panel text (“gain N×”) reports the fold increase in tumour selectivity "
 "(OV8/IOSE) of the most-selective variant relative to the consensus. The twelve PWMs shown are those that "
 "(i) retained at least eight quantified variants after quality control, (ii) contained a genuine "
 "high-affinity consensus in the screened set (highest-affinity surviving variant ≥ 0.8 of the PWM "
 "maximum), and (iii) yielded a non-consensus variant at least two-fold more tumour-selective than the "
 "consensus; among these, the twelve with the highest absolute tumour selectivity are displayed. This "
 "figure is illustrative of single motifs; all quantitative, library-wide statements are supported by "
 "Figures 2–6."),
("Figure 2. Distribution of tumour-versus-normal selectivity across all variants.",
 "Histogram of the log2 selectivity ratio, log2(OV8/IOSE), for all 37,832 quantified enhancer variants "
 "(grey bars). The orange curve is a Gaussian kernel-density estimate of the same distribution, scaled to "
 "the histogram counts. The vertical crimson line marks the empirical neutral centre, defined as the mode "
 "of the kernel-density estimate; the dotted black line marks a log2 ratio of 0 (equal activity in the two "
 "cell types, i.e. a raw OV8/IOSE ratio of 1). Positive values denote tumour-selective variants (greater "
 "activity in OV8) and negative values denote normal-selective variants (greater activity in IOSE). "
 "Selectivity is computed from the raw activity ratio without a pseudocount, so values correspond exactly "
 "to the OV8/IOSE ratios reported in the interactive dashboard."),
("Figure 3. Decomposition of variant-driven selectivity gains, by performance tier.",
 "Each of the six panels shows one performance tier. Tiers are defined by ranking all 5,683 PWMs by the "
 "absolute tumour selectivity of their best variant and taking, cumulatively, the top 0.1%, 1%, 10%, 20%, "
 "40%, and all PWMs (panel titles give the tier and the number of PWMs, n). Within a panel each point is "
 "one PWM, positioned by the change in tumour activity (x-axis) and the change in normal activity (y-axis) "
 "on moving from that PWM's consensus variant to its most tumour-selective variant; both changes are "
 "expressed in log2 units (best minus consensus) and clipped to ±3.2 for display. The dashed diagonal "
 "(y = x) marks selectivity-neutral changes, for which tumour and normal activity change by the same "
 "amount; points below the diagonal therefore represent a net gain in selectivity. Grey reference lines "
 "mark zero change on each axis, and the boxed text reports the median Δtumour and median Δnormal "
 "for that tier."),
("Figure 4. Median tumour and normal activity from consensus to best variant, by performance tier.",
 "Slope plot of population-median reporter activity for each of the six performance tiers (defined as in "
 "Figure 3; panel titles give the tier and n). In each panel the left-hand points are the median activity "
 "of the consensus variants and the right-hand points the median activity of the most tumour-selective "
 "variants, shown separately for tumour cells (OV8, red) and normal cells (IOSE, blue); the connecting "
 "lines indicate the direction and magnitude of each shift and the point labels give the median values "
 "(mean RNA/DNA ratio). The y-axis is scaled independently per panel so that tiers spanning very different "
 "absolute activity levels are each resolved."),
("Figure 5. Affinity of the most tumour-selective variant within each motif.",
 "Histogram of the within-PWM affinity rank of the most tumour-selective variant, computed for every PWM "
 "that retained at least four surviving variants with tumour activity ≥ 0.1. For each such PWM the most "
 "tumour-selective variant (highest OV8/IOSE) is identified and its relative-affinity rank among that PWM's "
 "variants is expressed as a percentile, where 0 denotes the weakest motif match in the set and 100 denotes "
 "the consensus. Bars are coloured by affinity band — the weakest fifth (red), the middle three-fifths "
 "(grey), and the strongest fifth, near the consensus (green) — and the dotted vertical lines mark the 20th "
 "and 80th-percentile band boundaries. The legend gives the percentage of PWMs whose most-selective variant "
 "falls in each band."),
("Figure 6. Relationship between motif affinity and selectivity across all variants.",
 "Two-dimensional density, by hexagonal binning, of all 37,832 variants according to their relative motif "
 "affinity (x-axis; 1.0 = consensus) and their log2 selectivity, log2(OV8/IOSE) (y-axis). Hexagon colour "
 "encodes the number of variants per bin on a linear scale (magma). The cyan line traces the median "
 "selectivity within twenty equal-width affinity bins, and the dotted white line marks the overall median "
 "selectivity across all variants."),
("Figure 7. The sequence-variant dimension is necessary to recover the most tumour-selective enhancers.",
 "A counterfactual comparison of two enhancer-library designs built from the same set of 5,683 motifs: a "
 "full variant library, which screens many binding-site variants per motif and may use the best variant of "
 "each, versus a consensus-only library, which contains only the single consensus (highest-affinity) "
 "binding site of each motif. (A) For each library, the tumour selectivity (OV8/IOSE) of the N-th most "
 "selective enhancer it can assemble — one enhancer per motif — is plotted against the number of enhancers "
 "selected, N (both axes logarithmic). The red curve is the full variant library (best surviving variant "
 "per motif, requiring OV8 activity ≥ 0.1); the grey curve is the consensus-only library (each motif's "
 "consensus variant, likewise requiring OV8 activity ≥ 0.1); the shaded band is the selectivity forfeited "
 "by restricting the design to consensus sequences. Dotted horizontal lines mark the 2-fold and 4-fold "
 "selectivity levels, annotated with the number of enhancers each library reaches at or above that level. "
 "(B) Number of motifs that yield a tumour-selective enhancer at or above a 2-, 4-, and 8-fold selectivity "
 "threshold for the full variant library (red) versus the consensus-only library (grey), on a logarithmic "
 "count axis; the bold label above each pair is the percentage of selective enhancers that the "
 "consensus-only design fails to recover. The two libraries contain essentially the same number of "
 "tumour-active motifs (5,683 vs 5,681), so the difference reflects selectivity, not coverage."),
]

METHODS = [
("Input data and core definitions.",
 "All figures are generated from a single per-variant table (one row per quantified synthetic enhancer "
 "variant). For each variant the table records its source PWM, transcription factor, binding-site sequence, "
 "number of tandem repeats, relative motif affinity, and mean reporter activity (RNA-to-DNA barcode ratio) "
 "in OV8 (tumour) and IOSE (normal) cells. Selectivity is defined as log2(OV8/IOSE), computed from raw "
 "activities without a pseudocount (no variant had zero activity in either cell type; the minimum observed "
 "activity was ~0.03), so the figures match the OV8/IOSE values shown in the dashboard. Relative motif "
 "affinity is the fraction-of-consensus score of each variant's binding site against its own PWM "
 "(∏_j p[j, seq_j] / max_b p[j, b]; 1.0 = consensus); it is pre-computed in the data table. For every "
 "PWM the consensus variant is the surviving variant of highest relative affinity, and the best variant is "
 "the most tumour-selective surviving variant (highest OV8/IOSE) among those with OV8 activity ≥ 0.1. "
 "Performance tiers are obtained by ranking all PWMs by their best variant's absolute selectivity and "
 "taking the cumulative top 0.1%, 1%, 10%, 20%, 40%, and all PWMs."),
("Per-figure plotting.",
 "Figure 1 plots, for twelve selected PWMs (selection criteria given in the legend), every variant in "
 "OV8-versus-IOSE activity space, coloured by relative affinity, with the consensus marked by a "
 "thick-outlined dot and the most-selective variant by a star; axes are shared and a y = x diagonal is "
 "drawn. Figure 2 histograms log2(OV8/IOSE) over all variants, overlays a Gaussian kernel-density estimate, "
 "and marks the density mode as the empirical neutral centre. Figure 3 plots, per tier, one point per PWM at "
 "(Δtumour, Δnormal) = (log2 best/consensus OV8, log2 best/consensus IOSE), with a y = x diagonal "
 "and the per-tier medians annotated. Figure 4 plots, per tier, the population-median OV8 and IOSE activity "
 "of consensus versus best variants as a two-point slope plot. Figure 5 histograms, over all qualifying "
 "PWMs, the within-PWM affinity percentile of the most-selective variant, coloured by affinity band. "
 "Figure 6 shows a linear-count hexagonal-bin density of all variants in affinity-versus-selectivity space "
 "with a binned median-selectivity trace. Figure 7 derives two enhancer-selection frontiers from the "
 "per-PWM summary: for the full variant library the per-PWM best-variant selectivities are sorted "
 "descending, and for the consensus-only library the per-PWM consensus-variant selectivities (consensus "
 "OV8 ≥ 0.1) are sorted descending; panel A plots both frontiers against selection rank with the gap "
 "shaded, and panel B counts, for each library, the PWMs whose selectivity meets each fold threshold."),
("Software.",
 "Figures were produced by a single Python script (make_figures.py) that reads the one data table "
 "(trend_ovca_variants.csv). Numerical work used numpy, pandas, and scipy (Gaussian kernel-density "
 "estimation). Figures 2, 3, 5, and 6 were drawn with matplotlib; Figures 1 and 4 were drawn with plotnine "
 "(a grammar-of-graphics library implementing the ggplot2 model). Each figure is written as a vector PDF "
 "for publication and a PNG for preview. The script and data table are provided as source code and source "
 "data for exact reproduction."),
]

CONCLUSIONS = [
("Figure 1.",
 "For motifs that possess a genuine high-affinity consensus, a non-consensus sequence variant is "
 "frequently the most tumour-selective element, and is often several-fold more selective than the "
 "consensus itself. Recovering the most cell-type-selective enhancers therefore requires sampling many "
 "sequence variants of a motif rather than testing the consensus alone."),
("Figure 2.",
 "The genome-wide selectivity distribution is centred slightly below an OV8/IOSE ratio of 1 — the neutral "
 "point is determined empirically rather than assumed — and is strongly asymmetric, containing far more "
 "tumour-selective than normal-selective variants. The platform thus preferentially yields tumour-selective "
 "regulatory elements, and selectivity calls should be referenced to the empirical centre."),
("Figure 3.",
 "Variant-driven selectivity gains arise predominantly from a reduction of activity in normal cells rather "
 "than an increase in tumour cells. A tumour-activity-raising component is also present and grows in the "
 "most selective tiers, so that the largest selectivity gains generally require both levers (lower normal "
 "and higher tumour activity) acting together."),
("Figure 4.",
 "Viewed as activity shifts, the same pattern is evident: across the most selective tiers, switching from "
 "the consensus to the best variant chiefly suppresses normal-cell activity while preserving — or, in the "
 "top tiers, modestly increasing — tumour-cell activity. This is the activity signature of variants that "
 "raise the threshold for activation in low-input (normal) cells while remaining active in high-input "
 "(tumour) cells."),
("Figure 5.",
 "There is no universal optimal binding affinity: depending on the motif, the most tumour-selective variant "
 "sits at low, intermediate, or near-consensus affinity, with only a mild skew toward lower affinity. This "
 "“Goldilocks” behaviour — an optimum that differs from motif to motif — is a direct rationale for "
 "screening a diverse panel of sequence variants per motif rather than relying on the consensus."),
("Figure 6.",
 "Across the full library, motif affinity does not predict selectivity: the median selectivity is flat with "
 "respect to affinity, and highly selective variants occur at every affinity level. Cell-type selectivity "
 "therefore cannot be inferred from binding-affinity scores alone and must be measured empirically, which is "
 "the function the TREND variant library provides."),
("Figure 7.",
 "The sequence-variant dimension of the library is necessary, not incidental: a hypothetical consensus-only "
 "design would forfeit most of the platform's tumour selectivity. Although consensus binding sites are "
 "themselves tumour-active for nearly every motif, restricting the design to consensus sequences fails to "
 "recover roughly 60% of the enhancers reaching ≥4-fold tumour selectivity (and a comparable fraction at "
 "≥2- and ≥8-fold). Because the most selective element of a motif is usually a non-consensus variant, "
 "systematically screening many variants per motif — the defining feature of TREND — is required to obtain "
 "the most cell-type-selective enhancers."),
]

# ----------------------------------------------------------------- builders
def styled(doc):
    n = doc.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(11)

def legend_doc():
    doc = Document(); styled(doc)
    doc.add_heading("TREND — Figure legends and figure-generation methods", level=0)
    doc.add_heading("Figure legends", level=1)
    for title, body in LEGENDS:
        p = doc.add_paragraph(); p.add_run(title).bold = True
        p.add_run(" " + body)
        p.paragraph_format.space_after = Pt(10)
    doc.add_heading("Figure generation (how each figure is plotted)", level=1)
    for title, body in METHODS:
        p = doc.add_paragraph(); p.add_run(title).bold = True
        p.add_run(" " + body)
        p.paragraph_format.space_after = Pt(10)
    out = HERE / "TREND_figure_legends.docx"; doc.save(out); return out

def conclusion_doc():
    doc = Document(); styled(doc)
    doc.add_heading("TREND — Conclusions by figure", level=0)
    doc.add_paragraph("Each entry states the conclusion supported by the corresponding figure. "
                      "These are interpretations and are deliberately kept separate from the figure legends.")
    for title, body in CONCLUSIONS:
        p = doc.add_paragraph(); p.add_run(title).bold = True
        p.add_run(" " + body)
        p.paragraph_format.space_after = Pt(10)
    out = HERE / "TREND_figure_conclusions.docx"; doc.save(out); return out

print("wrote", legend_doc())
print("wrote", conclusion_doc())
