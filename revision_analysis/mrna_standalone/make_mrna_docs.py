#!/usr/bin/env python
"""Write two Word documents for the TF-mRNA-vs-activity figures:
   1. TREND_mRNA_figure_legends.docx      — formal, manuscript-ready legends
   2. TREND_mRNA_figure_descriptions.docx — plain-language descriptions (author use)
Dep: python-docx  (pip install python-docx)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent

# --------------------------------------------------------------- formal legends
LEGENDS = [
("Figure 0. TF mRNA abundance versus enhancer selectivity (reference).",
 "Each point is one transcription factor (n = 755 expressed factors with sensors in "
 "the library), positioned by its differential mRNA expression between OV8 (tumour) "
 "and IOSE (normal) cells (log10 of the OV8/IOSE TPM ratio) and the median tumour "
 "selectivity (OV8/IOSE reporter activity) of its synthetic-enhancer sensors. The "
 "dotted line marks equal tumour/normal activity. The two quantities are uncorrelated "
 "(Pearson r = −0.02)."),
("Figure A. Selecting enhancers by TF mRNA performs no better than chance.",
 "Cumulative recovery of the 50 most tumour-selective transcription factors (ranked by "
 "their best sensor in the screen) as transcription factors are selected one at a time. "
 "Factors were ordered by descending differential mRNA expression (red), by the screen's "
 "measured selectivity (green; the ideal), or at random (grey dashed; analytical "
 "expectation). Selection by differential mRNA tracks random selection — the top 50 "
 "factors by mRNA recover 0 of the 50 most-selective enhancers."),
("Figure B. The determinant of enhancer selectivity lies within, not between, "
 "transcription factors.",
 "Per-sensor tumour selectivity (log2 OV8/IOSE; light points) for every sensor, plotted "
 "against its transcription factor's differential mRNA expression; black points are "
 "per-TF medians. The variance of selectivity within transcription factors (mean 0.34) "
 "exceeds the variance between TF medians (0.18) by ~1.9-fold; the dominant axis of "
 "variation is therefore between sensors of the same factor — a dimension inaccessible "
 "to any single per-TF measurement such as mRNA abundance."),
("Figure C. Differential mRNA does not classify transcription factors that yield "
 "selective enhancers.",
 "Receiver-operating-characteristic curve for differential mRNA expression used to "
 "classify transcription factors that have at least one ≥2-fold tumour-selective sensor "
 "(40% of factors). The area under the curve (0.54) is close to the chance value of "
 "0.50."),
("Figure D. Neither mRNA abundance nor prior literature identifies the best enhancers.",
 "Each point is a transcription factor positioned by its differential mRNA expression "
 "(x) and the selectivity of its most tumour-selective sensor (y); dotted lines mark "
 "2-fold mRNA over-expression and 2-fold sensor selectivity. Colour highlights three "
 "reference groups: the transcription factors with the most selective sensors in the "
 "screen (green), field-acknowledged ovarian-cancer factors (blue; PAX8, WT1, SOX17, "
 "MECOM, FOXM1), and the most mRNA-over-expressed factors (red). The most selective "
 "enhancers derive from factors that are neither the most over-expressed nor the "
 "canonical ovarian-cancer factors."),
("Figure E. The most over-expressed transcription factors are unremarkable enhancer "
 "sources.",
 "Distribution of best-sensor selectivity (log2 OV8/IOSE) for all transcription factors "
 "(grey) and for the 20 most mRNA-over-expressed factors (red), shown as area-normalized "
 "densities; vertical lines mark group medians (1.76-fold vs 1.87-fold). The most "
 "over-expressed factors are indistinguishable from the library as a whole in the "
 "selectivity of the enhancers they yield."),
("Figure F. The mRNA ranking and the enhancer ranking of transcription factors are "
 "unrelated.",
 "Transcription factors ranked by differential mRNA expression (x; rank 1 = most "
 "over-expressed) versus by best-sensor selectivity (y; rank 1 = most selective); the "
 "two rankings are uncorrelated (Spearman ρ = 0.10). Reference groups are coloured as in "
 "Figure D. None of the three groups aligns the two rankings."),
]

# --------------------------------------------------------------- plain descriptions
DESCRIPTIONS = [
("Shared terms.",
 "A 'sensor' is one synthetic enhancer (a TF binding-site variant in a reporter); each "
 "TF has many sensors (median 17). 'Sensor selectivity' is its OV8/IOSE activity ratio "
 "(log2: 0 = equal, 1 = 2-fold tumour-selective). 'Best-sensor selectivity' of a TF is "
 "its single most tumour-selective sensor — the best enhancer you could get by screening "
 "that TF. 'Differential TF expression' is log10(OV8 TPM / IOSE TPM): how much more the "
 "TF's mRNA is expressed in tumour than normal cells (the cheap RNA-seq metric). The "
 "whole point: the mRNA number does not tell you the best-sensor number."),
("Figure 0 (reference — the current view).",
 "The existing scatter: per-TF differential mRNA vs the median selectivity of that TF's "
 "sensors. It is flat (r = −0.02), i.e. 'no correlation'. Figures A–F reframe this null "
 "result as a positive, decision-relevant demonstration."),
("Figure A — recovery curve (how the lines are made).",
 "Pretend you must pick TFs to build enhancers, ranked by one criterion. The 'right "
 "answers' are the 50 TFs whose best sensor is most selective. X = how many TFs you have "
 "picked; Y = what % of those 50 you have captured. Red = pick in mRNA order; grey "
 "dashed = pick randomly (a straight diagonal); green = pick by the screen (the 50 best "
 "are your first 50, so it jumps to 100%). The red line sits on the grey line — mRNA "
 "picking is no better than blind; the top 50 by mRNA capture 0 of the 50 best."),
("Figure B — within- vs between-TF variance (the mechanism).",
 "mRNA gives one number per TF, so it can only ever explain how TFs differ from each "
 "other ('between-TF' variance). But selectivity varies far more among sensors of the "
 "SAME TF ('within-TF' variance = 1.9× larger), and a single per-TF number is blind to "
 "that. Each light point is one sensor (x = its TF's mRNA, y = its own selectivity); the "
 "huge vertical spread at every x is the within-TF variation; the black per-TF medians "
 "are flat. So even a perfect TF-ranker would miss the larger part of the signal."),
("Figure C — ROC / AUC (what the number means).",
 "Treat it as a yes/no test: does a TF have at least one good (≥2-fold selective) sensor? "
 "Use mRNA as the predictor and slide a cutoff; at each cutoff you catch some good TFs "
 "(true positives) but wrongly flag some bad ones (false positives). The ROC curve plots "
 "that trade-off. AUC = the chance that a random good TF has higher mRNA than a random "
 "bad TF: 1.0 = perfect, 0.5 = coin flip. We get 0.54 — essentially a coin flip."),
("Figure D — the TF map (why these dots are coloured).",
 "X = TF mRNA differential; Y = that TF's best-sensor selectivity. Three groups are "
 "highlighted with a legend: green = the screen's most selective TFs (e.g. the E2F "
 "family, MYC); blue = field-acknowledged ovarian-cancer TFs (PAX8, WT1, SOX17, MECOM, "
 "FOXM1); red = the most mRNA-over-expressed TFs. The top of the plot — the best "
 "enhancers — is occupied only by the green screen-winners. PAX8 (canonical, strongly "
 "over-expressed) sits low; so do the other literature TFs and the mRNA-top TFs. Leader "
 "lines connect labels to dots only where dots overlap."),
("Figure E — distributions (what 'density' and 'best-sensor selectivity' mean).",
 "Best-sensor selectivity = each TF's most selective sensor (one number per TF). We "
 "compare all TFs (grey) with the 20 most mRNA-over-expressed TFs (red). 'Density' is a "
 "histogram rescaled so each group sums to the same area — needed because one group has "
 "755 TFs and the other 20 — so it reads as 'what fraction of this group has this "
 "selectivity'. The two distributions overlap (medians 1.76× vs 1.87×): the most "
 "over-expressed TFs are not better enhancer sources."),
("Figure F — rank vs rank (the scramble).",
 "Rank all TFs by mRNA differential (x) and by best-sensor selectivity (y). If mRNA "
 "predicted enhancer quality, points would fall on a diagonal; instead they fill the "
 "square (Spearman ρ = 0.10 ≈ no relationship). The coloured/labelled points are the "
 "same three groups as Figure D: green (screen-best) line the bottom but spread across "
 "mRNA rank; red (mRNA-top) line the left but spread up the selectivity axis."),
]

def styled(doc):
    n = doc.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(11)

def build(title, items, fname):
    doc = Document(); styled(doc)
    doc.add_heading(title, level=0)
    for head, body in items:
        p = doc.add_paragraph(); p.add_run(head).bold = True
        p.add_run(" " + body); p.paragraph_format.space_after = Pt(10)
    out = HERE / fname; doc.save(out); return out

print("wrote", build("TREND — TF mRNA does not predict enhancer activity: figure legends",
                      LEGENDS, "TREND_mRNA_figure_legends.docx"))
print("wrote", build("TREND — TF mRNA vs enhancer activity: plain-language figure descriptions",
                      DESCRIPTIONS, "TREND_mRNA_figure_descriptions.docx"))
