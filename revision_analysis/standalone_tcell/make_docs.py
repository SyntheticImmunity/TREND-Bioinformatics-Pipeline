#!/usr/bin/env python
"""Generate two Word documents for the T-cell replication:
   1. TREND_tcell_figure_legends.docx     — detailed legends + a 'how each figure is plotted' section
   2. TREND_tcell_figure_conclusions.docx — the conclusion drawn from each figure (kept separate)
Run after make_figures.py. Dep: python-docx  (pip install python-docx)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent

# ----------------------------------------------------------------- content
LEGENDS = [
("Figure 1. Sequence-variant activity landscapes for representative transcription-factor motifs (T-cell activation).",
 "Each of the twelve panels corresponds to a single transcription-factor position-weight matrix (PWM) "
 "from the ENCODE/MotifDb motif collection; the panel title gives the transcription factor, the source "
 "database, and the motif accession. Within a panel, every plotted point is one synthetic enhancer "
 "variant — a tandem-repeat reporter built from a distinct binding-site sequence sampled from that PWM. "
 "Each variant is positioned by its measured activity in activated (stimulated) primary human T cells "
 "(x-axis) and in resting T cells (y-axis), expressed as the mean RNA-to-DNA barcode ratio (reporter "
 "activity) on linear axes shared across all panels. The fill colour of each point encodes the variant's "
 "relative binding affinity for its own PWM, computed as the fraction-of-consensus score (the product "
 "across motif positions of the base probability divided by that column's maximum probability; a value of "
 "1.0 corresponds to the consensus, i.e. optimal, sequence) and shown on the viridis colour scale at right. "
 "The dashed diagonal (y = x) denotes equal activity in activated and resting cells (no selectivity); "
 "points below the diagonal are activation-selective and points above are rest-selective. In every panel "
 "the consensus variant — operationally defined as the surviving variant with the highest relative affinity "
 "— is drawn as a thick black-outlined dot, and the most activation-selective variant — the variant with "
 "the highest Stim/Rest ratio among those with stim activity ≥ 0.1 — is drawn as a red star. The in-panel "
 "text (“gain N×”) reports the fold increase in activation selectivity (Stim/Rest) of the most-selective "
 "variant relative to the consensus. The twelve PWMs shown are those that (i) retained at least eight "
 "quantified variants after quality control, (ii) contained a genuine high-affinity consensus in the "
 "screened set (highest-affinity surviving variant ≥ 0.8 of the PWM maximum), and (iii) yielded a "
 "non-consensus variant at least two-fold more activation-selective than the consensus; among these, the "
 "twelve with the highest absolute activation selectivity are displayed. This figure is illustrative of "
 "single motifs; all quantitative, library-wide statements are supported by Figures 2–7."),
("Figure 2. Distribution of activation-versus-rest selectivity across all variants.",
 "Histogram of the log2 selectivity ratio, log2(Stim/Rest), for all 12,113 quantified enhancer variants "
 "(grey bars). The orange curve is a Gaussian kernel-density estimate of the same distribution, scaled to "
 "the histogram counts. The vertical crimson line marks the empirical neutral centre, defined as the mode "
 "of the kernel-density estimate; the dotted black line marks a log2 ratio of 0 (equal activity in the two "
 "states, i.e. a raw Stim/Rest ratio of 1). Positive values denote activation-selective variants (greater "
 "activity in stimulated cells) and negative values denote rest-selective variants. Selectivity is computed "
 "from the raw activity ratio without a pseudocount."),
("Figure 3. Decomposition of variant-driven selectivity gains, by performance tier.",
 "Each of the six panels shows one performance tier. Tiers are defined by ranking all 3,659 PWMs by the "
 "absolute activation selectivity of their best variant and taking, cumulatively, the top 0.1%, 1%, 10%, "
 "20%, 40%, and all PWMs (panel titles give the tier and the number of PWMs, n). Within a panel each point "
 "is one PWM, positioned by the change in stimulated activity (x-axis) and the change in resting activity "
 "(y-axis) on moving from that PWM's consensus variant to its most activation-selective variant; both "
 "changes are expressed in log2 units (best minus consensus) and clipped to ±3.2 for display. The dashed "
 "diagonal (y = x) marks selectivity-neutral changes; points below the diagonal therefore represent a net "
 "gain in selectivity. Grey reference lines mark zero change on each axis, and the boxed text reports the "
 "median Δstim and median Δrest for that tier."),
("Figure 4. Median stimulated and resting activity from consensus to best variant, by performance tier.",
 "Slope plot of population-median reporter activity for each of the six performance tiers (defined as in "
 "Figure 3; panel titles give the tier and n). In each panel the left-hand points are the median activity "
 "of the consensus variants and the right-hand points the median activity of the most activation-selective "
 "variants, shown separately for activated cells (Stim, red) and resting cells (Rest, blue); the connecting "
 "lines indicate the direction and magnitude of each shift and the point labels give the median values "
 "(mean RNA/DNA ratio). The y-axis is scaled independently per panel."),
("Figure 5. Affinity of the most activation-selective variant within each motif.",
 "Histogram of the within-PWM affinity rank of the most activation-selective variant, computed for every "
 "PWM that retained at least four surviving variants with stim activity ≥ 0.1. For each such PWM the most "
 "activation-selective variant (highest Stim/Rest) is identified and its relative-affinity rank among that "
 "PWM's variants is expressed as a percentile, where 0 denotes the weakest motif match in the set and 100 "
 "denotes the consensus. Bars are coloured by affinity band — the weakest fifth (red), the middle "
 "three-fifths (grey), and the strongest fifth, near the consensus (green) — and the dotted vertical lines "
 "mark the 20th and 80th-percentile band boundaries. The legend gives the percentage of PWMs whose "
 "most-selective variant falls in each band."),
("Figure 6. Relationship between motif affinity and selectivity across all variants.",
 "Two-dimensional density, by hexagonal binning, of all 12,113 variants according to their relative motif "
 "affinity (x-axis; 1.0 = consensus) and their log2 selectivity, log2(Stim/Rest) (y-axis). Hexagon colour "
 "encodes the number of variants per bin on a linear scale (magma). The cyan line traces the median "
 "selectivity within twenty equal-width affinity bins, and the dotted white line marks the overall median "
 "selectivity across all variants."),
("Figure 7. The sequence-variant dimension is necessary to recover the most activation-selective enhancers.",
 "A counterfactual comparison of two enhancer-library designs built from the same set of 3,659 motifs: a "
 "full variant library, which screens many binding-site variants per motif and may use the best variant of "
 "each, versus a consensus-only library, which contains only the single consensus (highest-affinity) "
 "binding site of each motif. (A) For each library, the activation selectivity (Stim/Rest) of the N-th most "
 "selective enhancer it can assemble — one enhancer per motif — is plotted against the number of enhancers "
 "selected, N (both axes logarithmic). The red curve is the full variant library (best surviving variant "
 "per motif, requiring stim activity ≥ 0.1); the grey curve is the consensus-only library (each motif's "
 "consensus variant, likewise requiring stim activity ≥ 0.1); the shaded band is the selectivity forfeited "
 "by restricting the design to consensus sequences. Dotted horizontal lines mark the 2-fold and 4-fold "
 "selectivity levels, annotated with the number of enhancers each library reaches. (B) Number of motifs "
 "that yield an activation-selective enhancer at or above a 2-, 4-, and 8-fold selectivity threshold for "
 "the full variant library (red) versus the consensus-only library (grey), on a logarithmic count axis; "
 "the bold label above each pair is the percentage of selective enhancers that the consensus-only design "
 "fails to recover. Both libraries contain the same number of stim-active motifs, so the difference "
 "reflects selectivity, not coverage."),
]

METHODS = [
("Input data and core definitions.",
 "All figures are generated from a single per-variant table (one row per quantified synthetic enhancer "
 "variant) built from the donor-1 activation-responsive enhancer screen. For each variant the table records "
 "its source PWM, transcription factor, binding-site sequence, number of tandem repeats, relative motif "
 "affinity, and mean reporter activity (RNA-to-DNA barcode ratio) in activated (Stim) and resting (Rest) "
 "T cells. Only variants quantified in both states are retained (12,113 variants across 3,659 PWMs). "
 "Selectivity is defined as log2(Stim/Rest), computed from raw activities without a pseudocount (no variant "
 "had zero activity in either state). Relative motif affinity is the fraction-of-consensus score of each "
 "variant's binding site against its own PWM (∏_j p[j, seq_j] / max_b p[j, b]; 1.0 = consensus), computed "
 "from the ENCODE/MotifDb PWM file with a 1e-3 column pseudocount and pre-computed into the data table. For "
 "every PWM the consensus variant is the surviving variant of highest relative affinity, and the best "
 "variant is the most activation-selective surviving variant (highest Stim/Rest) among those with stim "
 "activity ≥ 0.1. Performance tiers are obtained by ranking all PWMs by their best variant's absolute "
 "selectivity and taking the cumulative top 0.1%, 1%, 10%, 20%, 40%, and all PWMs. The analysis pipeline, "
 "definitions, and figure code are identical to the ovarian-cancer bundle; only the data (and the two cell "
 "states) differ, so the two figure sets are directly comparable."),
("Per-figure plotting.",
 "Figure 1 plots, for twelve selected PWMs (selection criteria given in the legend), every variant in "
 "Stim-versus-Rest activity space, coloured by relative affinity, with the consensus marked by a "
 "thick-outlined dot and the most-selective variant by a star; axes are shared and a y = x diagonal is "
 "drawn. Figure 2 histograms log2(Stim/Rest) over all variants, overlays a Gaussian kernel-density "
 "estimate, and marks the density mode as the empirical neutral centre. Figure 3 plots, per tier, one point "
 "per PWM at (Δstim, Δrest) = (log2 best/consensus Stim, log2 best/consensus Rest), with a y = x diagonal "
 "and the per-tier medians annotated. Figure 4 plots, per tier, the population-median Stim and Rest "
 "activity of consensus versus best variants as a two-point slope plot. Figure 5 histograms, over all "
 "qualifying PWMs, the within-PWM affinity percentile of the most-selective variant, coloured by affinity "
 "band. Figure 6 shows a linear-count hexagonal-bin density of all variants in affinity-versus-selectivity "
 "space with a binned median-selectivity trace. Figure 7 derives two enhancer-selection frontiers from the "
 "per-PWM summary: the full variant library sorts per-PWM best-variant selectivities descending, the "
 "consensus-only library sorts per-PWM consensus-variant selectivities descending; panel A plots both "
 "frontiers against selection rank with the gap shaded, and panel B counts, for each library, the PWMs "
 "whose selectivity meets each fold threshold."),
("Software.",
 "Figures were produced by a single Python script (make_figures.py) that reads the one data table "
 "(trend_tcell_variants.csv). Numerical work used numpy, pandas, and scipy (Gaussian kernel-density "
 "estimation). Figures 2, 3, 5, 6, and 7 were drawn with matplotlib; Figures 1 and 4 were drawn with "
 "plotnine (a grammar-of-graphics library implementing the ggplot2 model). Each figure is written as a "
 "vector PDF for publication and a PNG for preview."),
]

CONCLUSIONS = [
("Figure 1.",
 "As in ovarian cancer, for T-cell motifs that possess a genuine high-affinity consensus a non-consensus "
 "sequence variant is frequently the most activation-selective element, often several-fold more selective "
 "than the consensus itself. Recovering the most state-selective enhancers therefore requires sampling many "
 "sequence variants of a motif rather than testing the consensus alone."),
("Figure 2.",
 "The activation-selectivity distribution is centred slightly below a Stim/Rest ratio of 1 — the empirical "
 "neutral point (≈ −0.19 log2) is essentially identical to that measured in ovarian cancer — and is "
 "asymmetric toward activation-selective variants. The same empirically-anchored, asymmetric selectivity "
 "structure recurs in an unrelated cell system."),
("Figure 3.",
 "Variant-driven selectivity gains arise predominantly from a reduction of activity in resting cells rather "
 "than an increase in stimulated cells (median Δstim ≈ 0; median Δrest strongly negative in the top tiers), "
 "reproducing the ovarian-cancer result. Variants chiefly suppress the off-state, the activity signature of "
 "raising the threshold for activation in low-input (resting) cells while remaining active in high-input "
 "(stimulated) cells."),
("Figure 4.",
 "Viewed as activity shifts, the same pattern is evident: across the most selective tiers, switching from "
 "the consensus to the best variant chiefly suppresses resting-cell activity while preserving stimulated-cell "
 "activity — the cross-context analogue of the ovarian-cancer normal-suppression mechanism."),
("Figure 5.",
 "There is no universal optimal binding affinity in T cells either: depending on the motif, the most "
 "activation-selective variant sits at low, intermediate, or near-consensus affinity. This motif-specific "
 "“Goldilocks” behaviour — the same as in ovarian cancer — is a direct rationale for screening a diverse "
 "panel of sequence variants per motif rather than relying on the consensus."),
("Figure 6.",
 "Across the full library, motif affinity does not predict selectivity: the median selectivity is flat with "
 "respect to affinity, and highly selective variants occur at every affinity level — identical to the "
 "ovarian-cancer finding. Cell-state selectivity cannot be inferred from binding-affinity scores alone and "
 "must be measured empirically."),
("Figure 7.",
 "The sequence-variant dimension is necessary in T cells just as in ovarian cancer: a hypothetical "
 "consensus-only design would forfeit roughly 60% of the enhancers reaching ≥2-, ≥4-, and ≥8-fold "
 "activation selectivity, even though consensus sites are themselves stim-active for essentially every "
 "motif. Because the most selective element of a motif is usually a non-consensus variant, systematically "
 "screening many variants per motif is required to obtain the most state-selective enhancers — a "
 "design principle that holds across both screened cell systems."),
]

# ----------------------------------------------------------------- builders
def styled(doc):
    n = doc.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(11)

def legend_doc():
    doc = Document(); styled(doc)
    doc.add_heading("TREND (T-cell activation) — Figure legends and figure-generation methods", level=0)
    doc.add_heading("Figure legends", level=1)
    for title, body in LEGENDS:
        p = doc.add_paragraph(); p.add_run(title).bold = True
        p.add_run(" " + body)
        p.paragraph_format.space_after = Pt(10)
    doc.add_heading("Figure generation (how each figure is plotted)", level=1)
    for title, body in METHODS:
        p = doc.add_paragraph(); p.add_run(title).bold = True
        p.add_run(" " + body)
        p.paragraph_format.space_after = Pt(10)
    out = HERE / "TREND_tcell_figure_legends.docx"; doc.save(out); return out

def conclusion_doc():
    doc = Document(); styled(doc)
    doc.add_heading("TREND (T-cell activation) — Conclusions by figure", level=0)
    doc.add_paragraph("Each entry states the conclusion supported by the corresponding figure, with explicit "
                      "comparison to the ovarian-cancer bundle. These are interpretations and are deliberately "
                      "kept separate from the figure legends.")
    for title, body in CONCLUSIONS:
        p = doc.add_paragraph(); p.add_run(title).bold = True
        p.add_run(" " + body)
        p.paragraph_format.space_after = Pt(10)
    out = HERE / "TREND_tcell_figure_conclusions.docx"; doc.save(out); return out

print("wrote", legend_doc())
print("wrote", conclusion_doc())
